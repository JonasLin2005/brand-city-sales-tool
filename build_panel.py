#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Turn a contract book into one workbook per brand, ready for crawler.py.

Input: an .xlsx with one row per financing contract and three columns,
       brand, branch, funding date (header names are set below).
Output: <out>/<brand>.xlsx with one row per (province, branch city, month),
        column E holding the number of contracts placed that month, F and G
        left blank for crawler.py.

Branch names in the contract book carry the company name and a 分公司 suffix;
both are stripped, and the short name is looked up in branches.csv
(branch,province) to get its province.

USAGE
-----
  python build_panel.py --contracts sample_data/contracts_demo.xlsx \
      --out workbooks --strip-prefix 示例汽车服务有限公司 \
      --start 25年1月 --end 26年6月
"""

import argparse
import csv
import os
import re
import sys
from collections import defaultdict
from datetime import datetime

try:
    import openpyxl
    from openpyxl.styles import Font
except ImportError:
    sys.exit("Missing openpyxl. Run:  pip install -r requirements.txt")

SRC_BRAND, SRC_BRANCH, SRC_DATE = "品牌", "分公司", "放款日期"

HEADER = ["品牌名称", "展业省份", "展业城市（分公司）", "投放月份",
          "投放数（台）", "品牌城市月销量（台）", "月销量市占比"]
WIDTHS = {"A": 13.75, "B": 17.125, "C": 26.75, "D": 13.125,
          "E": 20.0, "F": 24.5, "G": 23.75}
FONT = Font(name="宋体", size=10)


def load_branches(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        return {r["branch"].strip(): r["province"].strip() for r in csv.DictReader(f)}


def short_branch(raw, prefix, suffix="分公司"):
    """'<company>广州第一分公司' -> '广州第一'; '南京市' -> '南京'."""
    if raw is None:
        return None
    s = str(raw).strip()
    if prefix and s.startswith(prefix):
        s = s[len(prefix):]
    if s.endswith(suffix):
        s = s[: -len(suffix)]
    if s.endswith("市") and len(s) > 2:
        s = s[:-1]
    return s or None


def parse_label(m):
    """'25年1月' -> (2025, 1)."""
    yy, rest = m.split("年")
    return 2000 + int(yy), int(rest.replace("月", ""))


def month_list(start, end):
    y, m = start
    out = []
    while (y, m) <= end:
        out.append((y, m))
        m += 1
        if m == 13:
            m, y = 1, y + 1
    return out


def pinyin_key():
    try:
        from pypinyin import lazy_pinyin
        return lambda s: "".join(lazy_pinyin(s))
    except ImportError:
        return lambda s: s


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--contracts", required=True)
    ap.add_argument("--out", default="workbooks")
    ap.add_argument("--branches", default=os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "branches.csv"))
    ap.add_argument("--strip-prefix", default="",
                    help="company name in front of every branch name")
    ap.add_argument("--start", default="25年1月")
    ap.add_argument("--end", default="26年6月")
    args = ap.parse_args()

    branch_prov = load_branches(args.branches)
    months = month_list(parse_label(args.start), parse_label(args.end))
    lo = datetime(*months[0], 1)
    hi_y, hi_m = months[-1]
    hi = datetime(hi_y + (hi_m == 12), hi_m % 12 + 1, 1)

    wb = openpyxl.load_workbook(args.contracts, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    ib, ic, idt = (header.index(h) for h in (SRC_BRAND, SRC_BRANCH, SRC_DATE))

    # brand -> branch -> (y, m) -> contracts
    data = defaultdict(lambda: defaultdict(lambda: defaultdict(int)))
    skipped = defaultdict(int)
    unmapped = set()
    for row in rows:
        brand, raw, when = row[ib], row[ic], row[idt]
        if brand is None:
            continue
        if not isinstance(when, datetime):
            skipped["no date"] += 1
            continue
        if not (lo <= when < hi):
            skipped["out of range"] += 1
            continue
        branch = short_branch(raw, args.strip_prefix)
        if branch is None:
            skipped["no branch"] += 1
            continue
        if branch not in branch_prov:
            unmapped.add(branch)
        data[str(brand).strip()][branch][(when.year, when.month)] += 1

    key = pinyin_key()
    os.makedirs(args.out, exist_ok=True)
    counted = sum(n for b in data.values() for c in b.values() for n in c.values())
    print(f"{counted} contracts counted, skipped: {dict(skipped) or 'none'}")
    if unmapped:
        print("Branches missing from branches.csv:", ", ".join(sorted(unmapped)))

    for brand, by_branch in sorted(data.items()):
        out = openpyxl.Workbook()
        ws = out.active
        ws.title = "表格视图"
        ws.append(HEADER)
        for col, w in WIDTHS.items():
            ws.column_dimensions[col].width = w
        branches = sorted(by_branch, key=lambda b: (key(branch_prov.get(b, "")), key(b)))
        for b in branches:
            for (y, m) in months:
                ws.append([brand, branch_prov.get(b, ""), b, f"{y - 2000}年{m}月",
                           by_branch[b].get((y, m), 0), None, None])
        for row in ws.iter_rows():
            for cell in row:
                cell.font = FONT
        for (cell,) in ws.iter_rows(min_row=2, min_col=7, max_col=7):
            cell.number_format = "0.00%"
        name = re.sub(r'[\\/:*?"<>|]', "_", brand) + ".xlsx"
        out.save(os.path.join(args.out, name))
        print(f"  {name}: {len(branches)} branches x {len(months)} months")


if __name__ == "__main__":
    main()
